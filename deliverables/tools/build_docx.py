# -*- coding: utf-8 -*-
"""Construit le dossier de projet OpsForge au format Word (.docx)."""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = r"C:\Users\DyllanTHOUVIGNON\Desktop\Work\opsforge"
DELIV = os.path.join(ROOT, "deliverables")
SHOTS = os.path.join(DELIV, "assets", "screenshots")
DIAG = os.path.join(DELIV, "assets", "diagrams")
FIGS = os.path.join(DELIV, "assets", "figures")
OUT = os.path.join(DELIV, "Dossier_de_projet_OpsForge_Dyllan_Thouvignon.docx")

# Figures derivees : recadrage des captures Prometheus sur leur zone utile
# (meme capture reelle, sans la zone vide ; les originales restent dans le depot).
os.makedirs(FIGS, exist_ok=True)
from PIL import Image as _Image
for _src, _dst, _h in (("06_prometheus_targets_up.png", "06_prometheus_targets_up_fig.png", 350),
                       ("07_prometheus_alert_firing.png", "07_prometheus_alert_firing_fig.png", 240)):
    _im = _Image.open(os.path.join(SHOTS, _src))
    _im.crop((0, 0, _im.width, _h)).save(os.path.join(FIGS, _dst))

TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_D = RGBColor(0x11, 0x5E, 0x59)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTE = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Segoe UI Semibold"
HEAD_FONT_L = "Segoe UI"
BODY_FONT = "Calibri"
CODE_FONT = "Consolas"

TEXT_W = 16.6  # cm utiles


# --------------------------------------------------------------------- outils
def no_proof(run):
    """Exclut le run du correcteur ET de la cesure automatique (Word)."""
    rPr = run._r.get_or_add_rPr()
    np = OxmlElement("w:noProof")
    rPr.append(np)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    rPr.append(lang)


def _rpr_fonts(el, name):
    rf = el.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def set_style_font(style, name, size, color=None, bold=None, italic=None):
    style.font.name = name
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "fr-FR")
    style.element.get_or_add_rPr().append(lang)


def shade_cell(cell, hexc):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexc)
    cell._tc.get_or_add_tcPr().append(shd)


def shade_par(par, hexc):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexc)
    par._p.get_or_add_pPr().append(shd)


def par_borders(par, left=None, bottom=None, top=None, space=4):
    pbdr = OxmlElement("w:pBdr")
    for edge, spec in (("top", top), ("left", left), ("bottom", bottom)):
        if not spec:
            continue
        color, sz = spec
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    par._p.get_or_add_pPr().append(pbdr)


def table_borders(table, color="D7DFE7", sz=4, inside=True):
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in edges:
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def cell_margins(table, top=70, start=110, bottom=70, end=110):
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement("w:" + k)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def fixed_layout(table):
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(lay)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:cantSplit")
    trPr.append(h)


# ------------------------------------------------------------- texte enrichi
TOKEN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)", re.S)


NBSP = "\u00a0"
APOS = "\u2019"


def fr_typo(t):
    """Typographie francaise : espace insecable avant la ponctuation haute,
    a l'interieur des guillemets, et apostrophe typographique."""
    t = t.replace("'", APOS)
    for sign in (":", ";", "!", "?", "»", "%"):
        t = t.replace(" " + sign, NBSP + sign)
    t = t.replace("« ", "«" + NBSP)
    return t


def add_rich(par, text, base_size=None, base_color=None, italic_all=False,
             _bold=False, _italic=False):
    """Rend **gras**, *italique* et `code`, y compris imbriques."""
    for chunk in TOKEN.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            add_rich(par, chunk[2:-2], base_size, base_color, italic_all,
                     _bold=True, _italic=_italic)
            continue
        if chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1])
            r.font.name = CODE_FONT
            _rpr_fonts(r._r, CODE_FONT)
            no_proof(r)
            r.font.size = Pt((base_size or 10.5) - 1.2)
            r.font.color.rgb = RGBColor(0x0B, 0x3F, 0x3B)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            add_rich(par, chunk[1:-1], base_size, base_color, italic_all,
                     _bold=_bold, _italic=True)
            continue
        else:
            r = par.add_run(fr_typo(chunk))
            if base_size:
                r.font.size = Pt(base_size)
            if base_color is not None:
                r.font.color.rgb = base_color
        r.bold = _bold or None
        r.italic = True if (_italic or italic_all) else None
    return par


# ------------------------------------------------------------------ document
doc = Document()

# --- page / marges
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.0)

# --- styles
st = doc.styles
normal = st["Normal"]
set_style_font(normal, BODY_FONT, 10.5, INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

h1 = st["Heading 1"]
set_style_font(h1, HEAD_FONT, 17, TEAL, bold=False)
h1.paragraph_format.space_before = Pt(2)
h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.keep_with_next = True
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

h2 = st["Heading 2"]
set_style_font(h2, HEAD_FONT, 12.5, TEAL_D, bold=False)
h2.paragraph_format.space_before = Pt(13)
h2.paragraph_format.space_after = Pt(5)
h2.paragraph_format.keep_with_next = True
h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

h3 = st["Heading 3"]
set_style_font(h3, HEAD_FONT, 11, INK, bold=False)
h3.paragraph_format.space_before = Pt(12)
h3.paragraph_format.space_after = Pt(3)
h3.paragraph_format.keep_with_next = True
h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

code_st = st.add_style("OF Code", WD_STYLE_TYPE.PARAGRAPH)
code_st.base_style = st["Normal"]
set_style_font(code_st, CODE_FONT, 8.6, RGBColor(0x0F, 0x17, 0x2A))
cpf = code_st.paragraph_format
cpf.space_before = Pt(2)
cpf.space_after = Pt(2)
cpf.line_spacing = 1.0
cpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
cpf.left_indent = Cm(0.35)
cpf.right_indent = Cm(0.1)
cpf.keep_together = True

cap_st = st.add_style("OF Caption", WD_STYLE_TYPE.PARAGRAPH)
cap_st.base_style = st["Normal"]
set_style_font(cap_st, HEAD_FONT_L, 8.6, MUTE, italic=True)
cap_st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap_st.paragraph_format.space_before = Pt(3)
cap_st.paragraph_format.space_after = Pt(12)
cap_st.paragraph_format.keep_together = True

part_st = st.add_style("OF PartTitle", WD_STYLE_TYPE.PARAGRAPH)
part_st.base_style = st["Normal"]
set_style_font(part_st, HEAD_FONT, 17, TEAL, bold=False)
part_st.paragraph_format.space_before = Pt(2)
part_st.paragraph_format.space_after = Pt(10)
part_st.paragraph_format.keep_with_next = True
part_st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

bul_st = st.add_style("OF Bullet", WD_STYLE_TYPE.PARAGRAPH)
bul_st.base_style = st["Normal"]
bul_st.paragraph_format.left_indent = Cm(0.6)
bul_st.paragraph_format.first_line_indent = Cm(-0.35)
bul_st.paragraph_format.space_after = Pt(4)

for _n, _sz, _ind, _af in (("TOC 1", 10.0, 0.0, 2), ("TOC 2", 9.5, 0.45, 1)):
    _s = st.add_style(_n, WD_STYLE_TYPE.PARAGRAPH)
    _s.base_style = st["Normal"]
    set_style_font(_s, BODY_FONT, _sz, INK)
    _s.paragraph_format.space_after = Pt(_af)
    _s.paragraph_format.space_before = Pt(0)
    _s.paragraph_format.line_spacing = 1.0
    _s.paragraph_format.left_indent = Cm(_ind)
    _s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

tbl_st = st.add_style("OF Table", WD_STYLE_TYPE.PARAGRAPH)
tbl_st.base_style = st["Normal"]
set_style_font(tbl_st, BODY_FONT, 9.4, INK)
tbl_st.paragraph_format.space_after = Pt(0)
tbl_st.paragraph_format.space_before = Pt(0)
tbl_st.paragraph_format.line_spacing = 1.06
tbl_st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


# ------------------------------------------------------------------ builders
def clear_outline(par):
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), "9")
    par._p.get_or_add_pPr().append(ol)


def P(text="", size=None, style=None, space_after=None, align=None, italic=False,
      keep_next=False, break_before=False, indent=None, color=None):
    p = doc.add_paragraph(style=style)
    if break_before:
        p.paragraph_format.page_break_before = True
    if text:
        add_rich(p, text, base_size=size, base_color=color, italic_all=italic)
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    if color is not None:
        for r in p.runs:
            r.font.color.rgb = color
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def H(level, text, break_before=False):
    p = doc.add_heading("", level=level)
    if break_before:
        p.paragraph_format.page_break_before = True
    add_rich(p, text)
    if level == 1:
        par_borders(p, bottom=("99C7C0", 8), space=6)
    return p


def BUL(items, size=None):
    for it in items:
        p = doc.add_paragraph(style="OF Bullet")
        r = p.add_run("— ")
        r.font.color.rgb = TEAL
        r.bold = True
        add_rich(p, it, base_size=size)
        if size:
            for r in p.runs[1:]:
                r.font.size = Pt(size)
    if items:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def CODE(lines, caption=None, keep_prev=True):
    if keep_prev and doc.paragraphs:
        prev = doc.paragraphs[-1]
        prev.paragraph_format.keep_with_next = True
        # Ces paragraphes se terminent par un chemin ou une commande insecable :
        # justifies, ils produisent des lignes anormalement etirees.
        prev.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body = lines.strip("\n").split("\n")
    ps = []
    for i, line in enumerate(body):
        p = doc.add_paragraph(style="OF Code")
        no_proof(p.add_run(line if line else " "))
        shade_par(p, "F4F7F8")
        spec_left = ("0F766E", 18)
        if len(body) == 1:
            par_borders(p, left=spec_left, top=("DDE5E8", 6), bottom=("DDE5E8", 6), space=3)
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(8)
        elif i == 0:
            par_borders(p, left=spec_left, top=("DDE5E8", 6), space=3)
            p.paragraph_format.space_before = Pt(7)
        elif i == len(body) - 1:
            par_borders(p, left=spec_left, bottom=("DDE5E8", 6), space=3)
            p.paragraph_format.space_after = Pt(8)
        else:
            par_borders(p, left=spec_left, space=3)
        if i < len(body) - 1:
            p.paragraph_format.keep_with_next = True
        ps.append(p)
    return ps


def TABLE(header, rows, widths, font=9.4, header_bg="0F766E", first_col_bold=False,
          band="F4F8F8", align_top=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fixed_layout(t)
    table_borders(t)
    cell_margins(t)

    def fill(cells, values, bold=False, white=False, size=font):
        for c, v, w in zip(cells, values, widths):
            c.width = Cm(w)
            para = c.paragraphs[0]
            para.style = doc.styles["OF Table"]
            add_rich(para, v, base_size=size)
            for r in para.runs:
                r.font.size = Pt(size)
                if bold:
                    r.bold = True
                if white:
                    r.font.color.rgb = WHITE
                    r.font.name = HEAD_FONT
                    _rpr_fonts(r._r, HEAD_FONT)
            if align_top:
                c.vertical_alignment = 0

    if header:
        row = t.add_row()
        fill(row.cells, header, bold=True, white=True, size=font - 0.2)
        for c in row.cells:
            shade_cell(c, header_bg)
        repeat_header(row)
        row_cant_split(row)
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True
    for i, values in enumerate(rows):
        row = t.add_row()
        row_cant_split(row)
        fill(row.cells, values)
        if first_col_bold:
            for r in row.cells[0].paragraphs[0].runs:
                r.bold = True
        if band and i % 2 == 1:
            for c in row.cells:
                shade_cell(c, band)
    P("", size=4, space_after=0)
    return t


def FIG(path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph(style="OF Caption")
    add_rich(c, caption)
    for r in c.runs:
        r.italic = True
        r.font.size = Pt(8.6)
        r.font.color.rgb = MUTE
    return c


def FIG2(p1, c1, p2, c2, width_cm):
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fixed_layout(t)
    no_borders(t)
    cell_margins(t, top=0, start=40, bottom=0, end=40)
    for j, (pth, cap) in enumerate(((p1, c1), (p2, c2))):
        cell = t.cell(0, j)
        cell.width = Cm(width_cm + 0.3)
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        par.add_run().add_picture(pth, width=Cm(width_cm))
        cc = t.cell(1, j)
        cc.width = Cm(width_cm + 0.3)
        cp = cc.paragraphs[0]
        cp.style = doc.styles["OF Caption"]
        cp.paragraph_format.space_after = Pt(6)
        add_rich(cp, cap)
        for r in cp.runs:
            r.italic = True
            r.font.size = Pt(8.6)
            r.font.color.rgb = MUTE
    row_cant_split(t.rows[0])
    P("", size=6, space_after=2)
    return t


def CALLOUT(text, bg="ECFDF5", border="5EEAD4"):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    fixed_layout(t)
    no_borders(t)
    cell_margins(t, top=110, start=160, bottom=110, end=160)
    row_cant_split(t.rows[0])
    c = t.cell(0, 0)
    c.width = Cm(TEXT_W)
    shade_cell(c, bg)
    p = c.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(0)
    add_rich(p, text)
    for r in p.runs:
        r.font.size = Pt(10)
    tcPr = c._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "18")
    el.set(qn("w:color"), "0F766E")
    bd.append(el)
    tcPr.append(bd)
    P("", size=6, space_after=4)
    return t


def add_toc(par):
    r = par.add_run()._r
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Sommaire — clic droit > Mettre à jour les champs."
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        r.append(e)


def add_page_field(par):
    r = par.add_run()._r
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = "PAGE"
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2):
        r.append(e)


# =========================================================== PAGE DE GARDE ==
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(58)
p.paragraph_format.space_after = Pt(16)
p.add_run().add_picture(os.path.join(DIAG, "logo_opsforge.png"), width=Cm(2.1))

p = P("DOSSIER DE PROJET", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, color=TEAL)
for r in p.runs:
    r.font.name = HEAD_FONT
    _rpr_fonts(r._r, HEAD_FONT)
    r.font.bold = True
p.runs[0].font.size = Pt(11)
pf = p.paragraph_format

p = P("OpsForge", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
r = p.runs[0]
r.font.name = HEAD_FONT
_rpr_fonts(r._r, HEAD_FONT)
r.font.size = Pt(40)
r.font.color.rgb = INK

p = P("Console locale de gestion d'incidents et chaîne DevOps de bout en bout",
      size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=22, color=MUTE)
for r in p.runs:
    r.font.name = HEAD_FONT_L
    _rpr_fonts(r._r, HEAD_FONT_L)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
par_borders(p, bottom=("0F766E", 12), space=1)

p = P("Titre professionnel **Administrateur système DevOps** — niveau 6",
      size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
p = P("Code titre TP-01414  ·  RNCP 36061", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
      space_after=30, color=MUTE)

info = [
    ("Candidat", "Dyllan Thouvignon"),
    ("Organisme de formation", "Liora (ex DataScientest)"),
    ("Session d'examen", "7 septembre 2026 — 09h30\nCampus Omnes Cœur Défense II, Courbevoie"),
    ("Type de projet", "Projet fil rouge indépendant, cahier des charges conçu par le candidat"),
    ("Dépôt Git", "ThDyllan/opsforge"),
    ("Candidat technique gelé", "branche phase6-operator-ux — commit a9ec694"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
fixed_layout(t)
no_borders(t)
cell_margins(t, top=60, start=0, bottom=60, end=120)
for k, v in info:
    row = t.add_row()
    row_cant_split(row)
    c0, c1 = row.cells
    c0.width, c1.width = Cm(5.4), Cm(9.6)
    p0 = c0.paragraphs[0]
    p0.style = doc.styles["OF Table"]
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(0)
    r = p0.add_run(fr_typo(k))
    r.font.size = Pt(9.6)
    r.font.color.rgb = MUTE
    r.font.name = HEAD_FONT_L
    _rpr_fonts(r._r, HEAD_FONT_L)
    p1 = c1.paragraphs[0]
    p1.style = doc.styles["OF Table"]
    p1.paragraph_format.space_after = Pt(0)
    for li, line in enumerate(v.split("\n")):
        if li:
            p1 = c1.add_paragraph(style="OF Table")
            p1.paragraph_format.space_after = Pt(0)
        rr = p1.add_run(fr_typo(line))
        rr.font.size = Pt(10.2)
        rr.bold = (li == 0)
        if "commit" in line or line.startswith("branche"):
            rr.bold = False

P("", size=8, space_after=0)
p = P("Version finale — août 2026", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTE)
p.paragraph_format.space_before = Pt(34)

# ============================================================ CORPS (sect 2) =
body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
body_sec.page_width, body_sec.page_height = Cm(21), Cm(29.7)
body_sec.top_margin, body_sec.bottom_margin = Cm(1.9), Cm(1.7)
body_sec.left_margin, body_sec.right_margin = Cm(2.2), Cm(2.2)
body_sec.header_distance, body_sec.footer_distance = Cm(1.1), Cm(1.0)

pg = OxmlElement("w:pgNumType")
pg.set(qn("w:start"), "1")
body_sec._sectPr.append(pg)

hdr = body_sec.header
hdr.is_linked_to_previous = False
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_after = Pt(2)
r = hp.add_run("OpsForge — Dossier de projet")
r.font.size = Pt(8.4)
r.font.color.rgb = MUTE
r.font.name = HEAD_FONT_L
_rpr_fonts(r._r, HEAD_FONT_L)
par_borders(hp, bottom=("DCE4E8", 6), space=3)

ftr = body_sec.footer
ftr.is_linked_to_previous = False
fp = ftr.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Dyllan Thouvignon  ·  TP-01414  ·  ")
r.font.size = Pt(8.4)
r.font.color.rgb = MUTE
r.font.name = HEAD_FONT_L
_rpr_fonts(r._r, HEAD_FONT_L)
add_page_field(fp)
for r in fp.runs:
    r.font.size = Pt(8.4)
    r.font.color.rgb = MUTE
    r.font.name = HEAD_FONT_L
    _rpr_fonts(r._r, HEAD_FONT_L)

# ------------------------------------------------------------------ SOMMAIRE
_som = doc.add_paragraph(style="OF PartTitle")
_som.add_run("Sommaire")
par_borders(_som, bottom=("99C7C0", 8), space=6)
tp = doc.add_paragraph()
tp.paragraph_format.space_after = Pt(0)
add_toc(tp)

# =================================================== LE PROJET EN 60 SECONDES
H(1, "Le projet en 60 secondes", break_before=True)

P("**OpsForge** est une console locale de gestion d'incidents : un opérateur y qualifie des signaux "
  "(alertes), ouvre et traite des incidents, applique des procédures contrôlées (runbooks) et conserve "
  "la preuve de chaque action dans un journal d'audit. Autour de cette application, j'ai construit et "
  "validé une chaîne DevOps locale complète : conteneurisation, intégration continue, "
  "sauvegarde/restauration, déploiement Kubernetes, supervision, et automatisation du déploiement "
  "de l'infrastructure.")

P("Les **trois compétences obligatoires** du titre sont couvertes, chacune avec une preuve distincte :",
  space_after=8)

TABLE(
    ["Compétence obligatoire", "Preuve principale", "Où"],
    [
        ["**Automatiser le déploiement d'une infrastructure**",
         "Playbook Ansible : toute l'infrastructure locale (cluster k3d, PostgreSQL, API, supervision) "
         "déployée et vérifiée en une commande idempotente — `/health` et `/ready` en 200, sinon échec",
         "§ 5.5\n§ 6"],
        ["**Gérer des containers**",
         "Image durcie non-root, orchestration Compose puis Kubernetes (securityContext complet, probes "
         "distinctes), stockage persistant dont la persistance est **prouvée** après destruction du pod",
         "§ 5.2\n§ 5.3\n§ 3.5"],
        ["**Exploiter une solution de supervision**",
         "Application instrumentée, Prometheus qui la scrape, dashboard Grafana, et une alerte réelle "
         "(`OpsForgeApiDown`) observée `inactive → pending → firing → résolue` lors d'une panne provoquée",
         "§ 5.1\n§ 3.6"],
    ],
    widths=[4.5, 10.3, 1.8],
)

P("Le périmètre est assumé : tout est local et démontrable ; il n'y a ni cloud, ni registre d'images, "
  "ni déploiement continu distant, ni authentification — ce sont des choix documentés, pas des oublis "
  "(§ 8).")

# ------------------------------------------------------------------- INTRO
H(1, "Introduction")

P("Je m'appelle Dyllan Thouvignon. Mon parcours était initialement orienté développement (BTS SIO "
  "option SLAM), puis j'ai suivi une formation DevOps en alternance. Mon activité en entreprise étant "
  "restée principalement orientée support technique, elle ne m'a pas fourni un projet DevOps "
  "suffisamment complet pour couvrir les attendus certificatifs.")

P("À l'issue de ma formation, j'ai donc conçu et réalisé **OpsForge**, un projet fil rouge indépendant, "
  "construit de juin à août 2026 pour préparer la certification. Ce n'est pas un projet de mon "
  "entreprise d'alternance : il n'utilise aucune donnée ni infrastructure professionnelle, et j'en ai "
  "défini le cahier des charges moi-même, à partir des compétences et des attendus du référentiel.")

P("Le fil conducteur du projet : chaque brique doit être **réellement démontrée** — tests exécutés, "
  "alerte réellement déclenchée, persistance réellement prouvée, déploiement réellement rejoué — et "
  "chaque limite est écrite noir sur blanc dans le dépôt. Ce dossier suit ce principe : il distingue "
  "systématiquement ce qui a été exécuté et vérifié de ce qui a seulement été conçu.")

# =============================================================== PARTIE 1 ==
H(1, "1. Compétences du référentiel couvertes par le projet", break_before=True)

H(2, "1.1 Les trois compétences obligatoires, critère par critère")
P("Le REAC (Référentiel Emploi Activités Compétences ; « CP » = compétence professionnelle) définit "
  "des critères de performance précis. Voici la confrontation exacte d'OpsForge à ces critères.")

H(3, "CP n° 2 — Automatiser le déploiement d'une infrastructure")
TABLE(
    ["Critère REAC", "Réponse OpsForge"],
    [
        ["« Les serveurs déployés sont fonctionnels »",
         "Le rôle Ansible `verify` échoue si un pod API n'est pas `Running` ou si `/health` et `/ready` "
         "ne répondent pas HTTP 200 (`/ready` exécute un `SELECT 1` sur PostgreSQL). Chaque étage est "
         "appliqué avec `wait: true`. Preuve rejouée le 12/08/2026 : `ok=22 changed=9 failed=0`, "
         "`/health → 200`, `/ready → 200`"],
        ["« L'architecture est conforme au cahier des charges »",
         "Le playbook orchestre les manifests versionnés de `k8s/` (namespaces, StatefulSet PostgreSQL "
         "+ PVC, Deployment API + NodePort, Prometheus + Grafana) — il déploie l'architecture "
         "documentée, sans rien redéfinir"],
        ["« Les scripts sont documentés »",
         "Cinq rôles courts et commentés, variables centralisées, `ansible/README.md`, justification et "
         "correspondance RNCP dans `docs/ANSIBLE.md`, décision d'architecture ADR 029"],
        ["Savoir associé : « outil de type Ansible ou Terraform »",
         "Ansible + collection `kubernetes.core` (choix argumenté en § 5.5)"],
    ],
    widths=[5.4, 11.2],
)
P("*Limite :* cible **k3d local**, pas un fournisseur cloud — l'activité-type « …dans le cloud » n'est "
  "pas revendiquée (la CP n° 4 relèvera de l'entretien technique).", size=9.8)

H(3, "CP n° 7 — Gérer des containers")
TABLE(
    ["Critère REAC", "Réponse OpsForge"],
    [
        ["« Les containers sont opérationnels »",
         "Compose : API + PostgreSQL avec healthchecks, démarrage de l'API conditionné à la santé de la "
         "base. Kubernetes : pods `1/1 Running` avec probes liveness/readiness distinctes (état constaté "
         "le 12/08/2026, § 3.5)"],
        ["« Les containers sont connectés au réseau »",
         "Réseau Compose interne (`db:5432`) ; Services Kubernetes ClusterIP (PostgreSQL interne "
         "uniquement) et NodePort 30080 exposé sur `127.0.0.1:8080`"],
        ["« Les containers sont connectés au stockage distant »",
         "PostgreSQL est connecté à un stockage persistant **externalisé du cycle de vie du conteneur** : "
         "volume nommé en Compose, PVC de 1 Gi monté par le StatefulSet en Kubernetes — la connexion du "
         "conteneur à un stockage porté par le système hôte figure parmi les savoir-faire associés à "
         "cette compétence dans le REAC. La persistance est **prouvée** : une donnée survit à la "
         "destruction puis à la recréation du pod (§ 3.5).\n"
         "*Limite assumée : dans l'environnement k3d utilisé, la StorageClass `local-path` reste locale "
         "au nœud — ce n'est pas un stockage réseau ni distribué ; une StorageClass réseau (NFS, CSI) la "
         "remplacerait sans modifier le manifest (§ 8).*"],
        ["« Les containers sont mis à jour »",
         "Cycle explicite : rebuild de l'image (tags `phase4` → `phase5` → `phase6`), `k3d image import`, "
         "rollout du Deployment ; `imagePullPolicy: Never` rend visible l'absence volontaire de registre"],
    ],
    widths=[5.4, 11.2],
)

H(3, "CP n° 10 — Exploiter une solution de supervision")
TABLE(
    ["Critère REAC", "Réponse OpsForge"],
    [
        ["« Les indicateurs définis sont pertinents »",
         "Métriques applicatives réelles (`opsforge_http_requests_total`, "
         "`opsforge_http_request_duration_seconds`, labels méthode / route template / code) et métrique "
         "`up` du scrape ; dashboard : disponibilité, volume, répartition par code, latence p95, "
         "répartition par route"],
        ["« Les alertes sont correctement interprétées »",
         "Règle `OpsForgeApiDown` (`up{job=\"opsforge-api\"} == 0` pendant 30 s). Panne provoquée, cycle "
         "complet observé et journalisé : `inactive → pending → firing → restauration → inactive` "
         "(validé en phase 5, **rejoué en direct le 12/08/2026** — § 3.6)"],
        ["« Les échanges avec les développeurs sont réguliers »",
         "*Limite du contexte individuel, assumée :* il n'y a pas d'équipe de développement — je tiens "
         "les deux rôles. La boucle supervision → développement existe réellement et elle est **tracée** "
         "dans les décisions du dépôt : ajout de `/ready` (ADR 022), labels de route en template pour "
         "maîtriser la cardinalité (ADR 017), configuration de scrape statique (ADR 018). En contexte "
         "d'équipe, ces constats seraient précisément le contenu des échanges avec les développeurs"],
    ],
    widths=[5.4, 11.2],
)

H(2, "1.2 Compétences partiellement mises en pratique")
P("Je ne revendique pas ces compétences comme couvertes ; le projet en met en œuvre une partie, "
  "exploitable lors des échanges avec le jury :")
TABLE(
    ["CP", "Mis en pratique", "Ce qui manque"],
    [
        ["CP 1 — Créer des serveurs par scripts",
         "Création scriptée et idempotente du nœud k3d ; scripts PowerShell et bash",
         "Pas de création de VM serveur au sens strict"],
        ["CP 3 — Sécuriser l'infrastructure",
         "Non-root, rootfs en lecture seule, capabilities supprimées, seccomp, secrets hors Git, scan "
         "Trivy, bind loopback",
         "Pas de référentiel ANSSI formalisé, ni pare-feu, ni TLS, ni authentification"],
        ["CP 5 — Environnement de test",
         "Base PostgreSQL éphémère par test, cluster k3d jetable, conteneur de service en CI",
         "Pas d'environnement mis à disposition d'une équipe"],
        ["CP 6 — Stockage des données",
         "PostgreSQL sur deux environnements, PVC prouvé, sauvegarde et restauration testées",
         "Pas de réplication ni de gestion formalisée des droits"],
        ["CP 8 — Mise en production avec une plateforme",
         "Déploiement Kubernetes automatisé par Ansible, vérifié de bout en bout",
         "Pas de pré-production et production distinctes, pas de CD"],
        ["CP 9 — Statistiques de services",
         "Indicateurs choisis et justifiés (ADR 017)",
         "Pas de SLA formalisés"],
    ],
    widths=[4.6, 6.6, 5.4],
)
P("**Non couvertes.** La CP n° 4 (mise en production dans le cloud) : aucun cloud n'a été déployé ; "
  "conformément aux modalités, elle relèvera du questionnement complémentaire de l'entretien technique. "
  "La CP n° 11 (échanger sur des réseaux professionnels, éventuellement en anglais) : elle est évaluée "
  "par une épreuve dédiée, le questionnaire professionnel en anglais — je note simplement que la "
  "documentation du dépôt et les messages de commit sont rédigés en anglais.")

# =============================================================== PARTIE 2 ==
H(1, "2. Cahier des charges", break_before=True)

H(2, "2.1 Contexte")
P("OpsForge est un projet fil rouge indépendant, réalisé dans le cadre de ma préparation au titre. "
  "J'en ai conçu le cahier des charges à partir du référentiel. Le domaine choisi — la gestion "
  "d'incidents — est inspiré de mon expérience de support : je connais le cycle « signal → prise en "
  "charge → procédure → traçabilité » pour l'avoir vécu au quotidien, ce qui m'a permis de définir un "
  "domaine métier crédible sans copier un outil existant ni utiliser la moindre donnée professionnelle "
  "réelle.")

H(2, "2.2 Problématique et objectifs")
P("**Problématique.** Quand un service supervisé se dégrade, comment garantir qu'un opérateur puisse "
  "qualifier le signal, décider d'une prise en charge, appliquer une procédure sûre et prouver ensuite "
  "chaque décision — et comment livrer cette application avec une vraie chaîne DevOps : tests "
  "automatisés, conteneurs, déploiement automatisé, supervision réelle, sauvegardes vérifiées ?")
P("**Objectifs.**", space_after=4)
BUL([
    "Une application métier fonctionnelle et démontrable, au cycle de vie strict et à l'audit systématique.",
    "Une chaîne DevOps locale couvrant les trois compétences obligatoires du titre.",
    "Une **preuve d'exécution réelle** pour chaque brique — pas seulement du code.",
    "Un ensemble simple, explicable et défendable à l'oral, aux limites connues et écrites.",
])

H(2, "2.3 Utilisateur et cas d'usage")
P("L'utilisateur est un **opérateur unique** (rôle : technicien d'exploitation). Parcours type, "
  "entièrement réalisable dans l'interface :", space_after=4)
BUL([
    "Un signal arrive sur un service du catalogue → il est enregistré comme **alerte** (`new`).",
    "L'opérateur **acquitte** l'alerte, puis décide s'il ouvre un **incident** (une alerte n'a qu'un "
    "incident actif à la fois).",
    "Il s'attribue l'incident, le passe en **investigation**, applique un **runbook** — checklist "
    "manuelle ou automatisation limitée à une liste d'actions approuvées dans le code.",
    "Chaque action alimente le **journal d'audit** et la timeline de l'incident.",
    "Il **résout** l'incident, puis résout l'alerte séparément : le cycle du signal et celui de la "
    "prise en charge sont indépendants.",
])
FIG(os.path.join(SHOTS, "01_overview.png"), 14.6,
    "**Figure 1** — Vue d'ensemble de la console au démarrage d'une prise de poste : incidents à "
    "traiter, alertes récentes avec leur état, état réel de la plateforme (à droite, « contrôles de la "
    "plateforme elle-même ») et services de démonstration.")

H(2, "2.4 Besoins")
P("**Fonctionnels.** Catalogue de services ; file d'alertes (cycle `new → acknowledged → resolved`, "
  "strictement en avant) ; incidents (cycle `open → investigating → resolved`, un seul incident actif "
  "par alerte, incident résolu en lecture seule) ; runbooks manuels à checklist et automatisations sur "
  "liste approuvée, jamais de commande arbitraire ; audit de toutes les mutations, y compris les échecs "
  "contrôlés ; console multipage utilisable sans outil externe ; distinction affichée entre contrôles "
  "réels de la plateforme et états métier simulés.")
P("**Techniques.** API FastAPI + PostgreSQL ; image Docker non-root avec healthcheck ; environnement "
  "Docker Compose complet ; CI à chaque push (lint, tests SQLite, test d'intégration PostgreSQL, build, "
  "scan de vulnérabilités) ; sauvegarde `pg_dump` scriptée avec restauration de vérification sans "
  "risque ; déploiement Kubernetes local (Deployment, StatefulSet, ConfigMap, Secret, PVC, probes, "
  "NodePort) ; supervision Prometheus/Grafana avec une alerte démontrable ; **déploiement complet de "
  "l'infrastructure en une commande idempotente** (besoin ajouté en cours de projet — § 2.7) ; "
  "traçabilité générale (Git, documentation par phase, décisions d'architecture).")

H(2, "2.5 Contraintes")
BUL([
    "Poste **Windows 11 + Docker Desktop** : ce choix contraint directement l'outillage (k3d plutôt "
    "qu'une VM, PowerShell pour les sauvegardes, control node Ansible conteneurisé).",
    "**Projet individuel** mené en parallèle de mon activité professionnelle, sur une période resserrée "
    "(juin → août 2026) : phases courtes, finies et validées une à une.",
    "**Aucune donnée réelle**, budget **zéro cloud**, et **aucune exécution de commande arbitraire** "
    "par l'application (vérifié par inspection et par test).",
    "Changement de poste de travail en cours de projet : la reproductibilité (Git, images, manifests) "
    "devait le permettre — et l'a permis.",
])

H(2, "2.6 Livrables et critères d'acceptation")
P("**Livrables.** Le dépôt Git complet (application, tests, Dockerfile, Compose, `k8s/`, `ansible/`, "
  "scripts, CI) ; la documentation projet (`docs/` : architecture, guides, 29 décisions d'architecture, "
  "risques, un fichier de vérification daté par phase) ; le présent dossier et son support de "
  "présentation.")
P("**Critères d'acceptation globaux.** Environnement Compose fonctionnel (`/health`, `/ready`, "
  "console) ; tests verts en local et en CI ; pods Kubernetes `1/1 Ready` et application joignable "
  "depuis Windows ; persistance PostgreSQL prouvée après recréation du pod ; alerte de supervision "
  "réellement déclenchée puis résolue ; déploiement d'infrastructure rejouable en une commande "
  "auto-vérifiée ; sauvegarde produite et restauration vérifiée sans toucher la base principale ; "
  "chaque phase validée explicitement et documentée.")

H(2, "2.7 Hors périmètre et évolution du périmètre")
P("**Exclusions volontaires** (décisions documentées — ADR, `docs/RISKS_AND_TECHNICAL_DEBT.md`) : "
  "cloud et Terraform ; registre d'images et déploiement continu distant ; authentification ; Alembic ; "
  "Alertmanager et notifications ; Helm ; React ; haute disponibilité.")
P("**Évolution du périmètre — l'histoire réelle.** Le projet est parti d'un MVP volontairement réduit "
  "(application + Compose + 7 tests), cadré par un document initial fixant six phases prévisionnelles "
  "et une règle : rien n'entre dans une phase sans décision explicite. Deux extensions ont été décidées "
  "en cours de route, tracées dans le dépôt :", space_after=4)
BUL([
    "**La phase 6 est devenue une phase produit** : le tableau de bord unique ne permettait pas une "
    "démonstration opérateur crédible → console multipage, Command Center par incident, runbooks "
    "managés, règles de domaine durcies, campagne de tests portée de 8 à 38 tests.",
    "**Ansible a été ajouté en fin de projet** : en confrontant le projet aux critères exacts du REAC, "
    "j'ai constaté que mon déploiement k3d, documenté mais **manuel**, ne prouvait pas la compétence "
    "obligatoire d'automatisation. J'ai fermé cet écart par un périmètre ciblé — automatiser le "
    "déploiement existant, sans rien redéfinir (ADR 029, § 5.5, § 6).",
])
P("Cette progression itérative est une caractéristique du projet : chaque phase validée fige un socle, "
  "et une relecture du référentiel a déclenché une correction de périmètre au bon moment.")

# =============================================================== PARTIE 3 ==
H(1, "3. Spécifications techniques", break_before=True)

H(2, "3.1 Architecture logique")
FIG(os.path.join(DIAG, "schema_architecture.png"), TEXT_W,
    "**Figure 2** — Architecture logique : l'application et sa supervision réelle (haut), puis les deux "
    "chaînes d'outillage (bas). La CI valide le code et l'image sans rien publier ; Ansible reconstruit "
    "l'image localement et déploie l'infrastructure.")
P("Les deux chaînes du bas sont **indépendantes** : la CI valide le code et l'image mais ne publie ni "
  "ne déploie rien (il n'y a pas de registre) ; Ansible reconstruit l'image localement et déploie "
  "l'infrastructure locale.")

H(2, "3.2 Topologie de déploiement")
FIG(os.path.join(DIAG, "schema_topologie.png"), TEXT_W,
    "**Figure 3** — Topologie déployée : tout tient sur un poste Windows 11 avec Docker Desktop. Le "
    "control node Ansible est un conteneur éphémère qui pilote le cluster k3d ; l'API est exposée sur "
    "`127.0.0.1:8080` via le NodePort 30080.")
P("L'accès à la supervision se fait par `kubectl port-forward` (Prometheus 9090, Grafana 3000) — un "
  "choix documenté qui évite d'exposer davantage un environnement local.")

H(2, "3.3 Le domaine applicatif")
P("Six objets : `Service`, `Alert`, `Incident`, `Runbook`, `RunbookExecution`, `AuditLog`. Les règles "
  "sont appliquées **côté serveur** et testées négativement :", space_after=4)
BUL([
    "transitions strictement en avant (alerte `new → acknowledged → resolved`, résolution directe "
    "possible ; incident `open → investigating → resolved`, sans réouverture) — une transition invalide "
    "renvoie HTTP 409 ;",
    "une alerte n'a qu'un **seul incident actif** (le doublon renvoie 409 avec l'identifiant de "
    "l'incident existant) ; l'incident hérite du service de son alerte source ; résoudre l'incident ne "
    "résout pas l'alerte ;",
    "**aucune exécution arbitraire** : 5 clés d'automatisation approuvées dans le code, clé inconnue → "
    "422 ; les runbooks définis dans le code sont « managés » (lecture seule, PATCH → 409) ; un runbook "
    "manuel ne peut pas être déclaré réussi avec une checklist incomplète ;",
    "toute mutation significative et toute tentative d'exécution — y compris les échecs contrôlés — "
    "produisent une entrée d'audit.",
])
P("**Structure du code :** `main.py` (démarrage, `/health`, `/ready`, `/metrics`, middleware de "
  "métriques), `api.py` (22 routes JSON et gardes), `web.py` (8 sections de console, 18 routes HTML), "
  "`domain.py` (transitions), `runbooks.py` (liste approuvée et moteur), `models.py` / `schemas.py`, "
  "`seed.py` (scénario de démonstration idempotent), `migrations.py` (pont additif de schéma au "
  "démarrage).")
FIG(os.path.join(SHOTS, "03_incident_command_center.png"), 11.8,
    "**Figure 4** — Le Command Center d'un incident : contexte opérationnel, alerte source, runbooks "
    "compatibles (manuel ou automatisé, niveau de risque), chronologie issue du journal d'audit et "
    "historique des exécutions.")

H(2, "3.4 Conteneurs : de Compose à Kubernetes")
BUL([
    "**Image API** : base `python:3.12-slim` épinglée par digest, utilisateur non-root UID 10001, "
    "`HEALTHCHECK` intégré (§ 5.2).",
    "**Compose** (développement, tests, sauvegardes) : PostgreSQL avec healthcheck `pg_isready` et API "
    "démarrée seulement quand la base est saine (`depends_on: service_healthy`) ; base publiée sur "
    "`127.0.0.1` uniquement ; `tests/` monté en lecture seule.",
    "**Kubernetes** (déploiement orchestré local) : mêmes conteneurs, contraintes renforcées — init "
    "container `wait-for-postgres`, probes distinctes (`/ready` = base joignable, `/health` = processus "
    "vivant), securityContext complet, requests et limits (§ 5.3).",
    "**Mise à jour** : rebuild → `k3d image import` → rollout ; pas de registre (`imagePullPolicy: "
    "Never`, ADR 015).",
])

H(2, "3.5 Stockage et persistance (prouvée)")
P("PostgreSQL utilise un **StatefulSet** (identité stable `postgres-0`, attachement stable au stockage) "
  "avec un **PVC de 1 Gi** en StorageClass `local-path` : le stockage est externalisé du cycle de vie "
  "du conteneur et porté par l'hôte. État constaté sur le cluster déployé par Ansible (12/08/2026, "
  "extrait de `evidence/kubernetes_state.txt`) :")
CODE("""pod/opsforge-api-68cbc5b54b-wrdqg   1/1     Running
pod/postgres-0                      1/1     Running
service/opsforge-api   NodePort    8000:30080/TCP
persistentvolumeclaim/postgres-data   Bound   1Gi   RWO   local-path""")
P("**Preuve de persistance** (rejouée le 12/08/2026, `evidence/pvc_persistence.txt`) : un "
  "marqueur est inséré, le pod est détruit, le StatefulSet le recrée — l'UID change, la donnée "
  "survit :")
CODE("""pod UID avant  : 554c3801-0b0e-42b7-8bf2-99ced94fd488
INSERT INTO persistence_check VALUES ('dossier-v2-persisted');
pod "postgres-0" deleted  ->  le StatefulSet recree le pod
pod UID apres  : e035ce38-526f-4b3b-9ef5-91aadb3c09f1   (pod reellement nouveau)
SELECT marker  : dossier-v2-persisted                    (donnee retrouvee)""")
P("*Limite :* `local-path` est local au nœud — la persistance couvre la recréation du pod, pas la "
  "destruction du cluster ; c'est le rôle des sauvegardes (§ 3.7).", size=9.8)

H(2, "3.6 Supervision : réel contre simulé")
CALLOUT("**Le point d'honnêteté central du projet.** Prometheus supervise **OpsForge lui-même** "
        "(l'API déployée dans k3d). Les statuts métier des services du catalogue sont des données de "
        "démonstration saisies dans l'application, et la console l'affiche explicitement.")
FIG(os.path.join(SHOTS, "05_monitoring.png"), 14.0,
    "**Figure 5** — La page Monitoring sépare la « Supervision réelle » (health, readiness, métriques, "
    "Prometheus/Grafana, règle testée) des « États métier simulés » (source : saisie OpsForge), et "
    "affiche la limite : les alertes Prometheus ne sont pas ingérées dans OpsForge.")
P("La chaîne réelle : middleware FastAPI → `/metrics` → scrape Prometheus toutes les 15 s (cible "
  "statique, job `opsforge-api`) → dashboard Grafana provisionné par ConfigMaps (5 panneaux : "
  "disponibilité, volume, codes, latence p95, routes) → règle `OpsForgeApiDown` (`up == 0` pendant "
  "30 s, calibrée sur l'intervalle de scrape).")
P("**Panne provoquée et alerte réelle** — cycle complet journalisé (12/08/2026, "
  "`evidence/prometheus_alert_cycle.txt`) :")
CODE("""$ kubectl -n opsforge scale deployment/opsforge-api --replicas=0
t+20s : up=0  OpsForgeApiDown=inactive
t+40s : up=0  OpsForgeApiDown=pending
t+70s : up=0  OpsForgeApiDown=firing      >>> FIRING observe <<<
$ kubectl -n opsforge scale deployment/opsforge-api --replicas=1
t+20s : up=1  OpsForgeApiDown=firing
t+40s : up=1  OpsForgeApiDown=inactive    >>> retabli, alerte resolue <<<""")
FIG(os.path.join(FIGS, "06_prometheus_targets_up_fig.png"), TEXT_W,
    "**Figure 6** — Prometheus scrape réellement l'API : la cible `opsforge-api (1/1 up)` pointe sur "
    "`/metrics` du Service Kubernetes, état `UP`.")
FIG(os.path.join(FIGS, "07_prometheus_alert_firing_fig.png"), TEXT_W,
    "**Figure 7** — La même instance pendant la panne provoquée : la règle `OpsForgeApiDown` est "
    "passée en `firing (1)`.")
FIG(os.path.join(SHOTS, "08_grafana_dashboard.png"), 14.4,
    "**Figure 8** — Le dashboard « OpsForge Monitoring » pendant la session de preuve : disponibilité "
    "UP, volume de requêtes (le creux correspond à la panne provoquée), 319 réponses 200 et 2 réponses "
    "503 au redémarrage, latence p95 et répartition par route.")
P("*Limites :* pas d'Alertmanager (l'alerte est visible dans Prometheus, elle n'est pas routée) ; "
  "stockage Prometheus et Grafana éphémère ; métriques techniques HTTP, pas métier.", size=9.8)

H(2, "3.7 Sauvegarde et restauration")
P("Deux scripts PowerShell ciblent le PostgreSQL de Compose : `backup.ps1` (`pg_dump --format=custom` "
  "exécuté dans le conteneur, archive copiée vers `backups/`, refus d'une archive vide) et "
  "`restore.ps1` (validation `pg_restore --list`, puis **par défaut** restauration dans une base "
  "temporaire de vérification ; la base principale exige `-MainDatabase` **et** la saisie exacte de "
  "`RESTORE`). Preuve rejouée le 12/08/2026 (`evidence/backup_restore.txt`) :")
CODE("""Backup created: backups\\opsforge_backup_20260812_133513.dump   (24 543 octets)
Restore verified in temporary database 'opsforge_restore_verify' (6 public tables).""")
P("*Limite :* sauvegarde locale, non planifiée, non chiffrée, sans rotation — un mécanisme "
  "démontrable, pas une stratégie d'entreprise.", size=9.8)

H(2, "3.8 Sécurité")
P("En couches, toutes documentées : (1) **par conception** — aucune exécution de commande arbitraire, "
  "liste d'automatisations approuvées, absence de `subprocess` et d'`eval` vérifiée ; (2) **conteneur** "
  "— non-root UID 10001, digest épinglé ; (3) **Kubernetes** — securityContext complet vérifié en "
  "conditions réelles (§ 5.3) ; (4) **secrets** — `.env` et `secret.local.yaml` ignorés par Git, Secret "
  "généré à la volée par Ansible ; (5) **chaîne d'approvisionnement** — scan Trivy **advisory** à "
  "chaque build : le dernier scan de l'image épinglée rapportait 19 HIGH et 3 CRITICAL d'origine "
  "Debian, sans correctif disponible. Une CI verte ne signifie donc pas « image sans vulnérabilité », "
  "et c'est assumé (ADR 013, § 5.4).")

H(2, "3.9 Environnements et versions")
TABLE(
    ["Composant", "Version"],
    [
        ["Application", "Python 3.12 (digest épinglé), FastAPI / SQLAlchemy 2.x, OpsForge v0.2.0"],
        ["Base de données", "PostgreSQL 16 (`postgres:16-alpine`) partout"],
        ["Cluster", "k3d v5.9.0 / k3s v1.35.x, kubectl client v1.34"],
        ["Supervision", "Prometheus v2.55.1, Grafana 11.3.1"],
        ["Control node Ansible",
         "ansible-core 2.17.14, kubernetes.core 6.5.0, client Python kubernetes 36.0.3, kubectl "
         "v1.31.5, k3d v5.9.0, CLI Docker 27.5.1 — versions validées, épinglées dans `ansible/Dockerfile`"],
        ["Intégration continue",
         "GitHub Actions (checkout@v5, setup-python@v6, trivy-action@v0.36.0)"],
        ["Poste de travail", "Windows 11 + Docker Desktop"],
    ],
    widths=[4.2, 12.4],
    first_col_bold=True,
)
P("*Note :* le kubectl du control node (v1.31.5) est plus ancien que le k3s du cluster (v1.35.x), "
  "au-delà de la fenêtre officielle de ±1 version mineure ; il est fonctionnel sur les opérations "
  "utilisées et validé tel quel — l'alignement est listé en évolution (§ 8).", size=9.8)

# =============================================================== PARTIE 4 ==
H(1, "4. Démarche de travail et outils", break_before=True)

H(2, "4.1 Le découpage en phases et leur état de validation")
TABLE(
    ["Phase", "Contenu", "Validation"],
    [
        ["1 — MVP local", "Application + Compose + 7 tests", "16/06/2026"],
        ["2 — Intégration continue", "GitHub Actions : tests, build, scan Trivy (run vert)",
         "18/06/2026"],
        ["3 — Sauvegarde et sécurité", "Sauvegarde/restauration, stratégie de secrets", "06/07/2026"],
        ["4 — Kubernetes",
         "Cluster k3d, PostgreSQL + PVC, API, **preuve de persistance**", "09/07/2026"],
        ["5 — Supervision",
         "`/metrics`, Prometheus, Grafana, **alerte réellement déclenchée**", "14/07/2026"],
        ["6 — Produit opérateur et preuves",
         "Console multipage, domaine durci, durcissement Kubernetes, **Ansible** (CP n° 2), puis revue "
         "manuelle : deux défauts réels corrigés et re-testés (38 tests)",
         "19/08/2026"],
    ],
    widths=[3.6, 10.6, 2.4],
    first_col_bold=True,
)
P("Chaque phase a un périmètre écrit, une **Definition of Done** vérifiable, un fichier de preuve daté "
  "(`docs/PHASE<i>_VERIFICATION.md`) et une validation explicite. Un protocole écrit en encadre l'avant "
  "et l'après : départ propre, périmètre figé — « aucune idée nouvelle n'entre silencieusement dans la "
  "phase en cours » — preuves consignées avant validation. Les choix importants sont consignés dans "
  "**29 décisions d'architecture**, qui servent aussi de préparation à l'oral.")

H(2, "4.2 Workflow Git")
P("Le workflow a évolué avec le projet : commits directs sur `main` pour les phases 1 à 5 (un commit "
  "de validation par phase), première branche dédiée pour le candidat produit de la phase 6, puis — "
  "pour l'audit final et l'ajout d'Ansible — un vrai cycle par branches et **Pull Requests avec commits "
  "de merge**, qui préserve l'historique des branches.")
P("Une branche d'expérimentation a servi de terrain de revue puis n'a **jamais été fusionnée** : une "
  "branche d'intégration propre a été ré-implémentée en six commits revus, intégrant les corrections "
  "identifiées (dont une vraie régression d'interface et la restauration du signal d'échec Trivy). Le "
  "même cycle a servi jusqu'au bout : les deux défauts de la revue manuelle finale ont été corrigés en "
  "branches dédiées depuis le candidat gelé, revalidés puis fusionnés (PR #4 et #5), la clôture étant "
  "enregistrée par une PR documentaire (PR #6). Le détail commit par commit est en annexe A.")

H(2, "4.3 Outils")
P("Python 3.12 / FastAPI / SQLAlchemy / Pydantic / Jinja2 · pytest (SQLite en mémoire et intégration "
  "PostgreSQL) · Ruff · Docker et Docker Compose · GitHub Actions · Trivy · k3d/k3s · kubectl · Ansible "
  "(`kubernetes.core`) · prometheus-client, Prometheus, Grafana · `pg_dump` / `pg_restore` en "
  "PowerShell · Git et GitHub.")

H(2, "4.4 Conduite du projet et usage de l'IA")
P("OpsForge est un **projet individuel** : pas de client, pas d'équipe — je tiens tous les rôles, et "
  "chaque validation de phase est ma décision, tracée dans le dépôt. Dans ce cadre, j'ai utilisé des "
  "assistants d'IA comme outils d'aide à la conception, à l'implémentation et à la revue, selon des "
  "règles écrites dans le dépôt (`docs/ENGINEERING_CHARTER.md`).")
P("Je suis resté responsable du cadrage, des choix techniques, des arbitrages et de la validation, et "
  "les changements importants ont été vérifiés et testés avant intégration. Ce processus explique le "
  "cycle de revue visible dans l'historique Git — « proposition → implémentation → revue → correction → "
  "validation » ; dans ce dossier, « revue » désigne cette revue outillée, conduite sous ma "
  "responsabilité.")

# =============================================================== PARTIE 5 ==
H(1, "5. Réalisations significatives", break_before=True)
P("Chaque réalisation suit le même fil : **besoin → extrait utile → choix → preuve → limite**. Les "
  "extraits proviennent du dépôt au commit final `a9ec694`, condensés pour la lecture (les coupures "
  "sont signalées par `# […]`) ; les fichiers complets sont dans Git.")

H(2, "5.1 Instrumentation et alerte de supervision (CP n° 10)")
P("**Besoin.** Superviser réellement OpsForge : des métriques produites par l'application elle-même, "
  "et une alerte qui se déclenche vraiment.")
CODE("""HTTP_REQUESTS_TOTAL = Counter(
    "opsforge_http_requests_total",
    "Total HTTP requests handled by OpsForge.",
    ["method", "route", "status_code"],
)
# [...] middleware : chronometre chaque requete et alimente compteur + histogramme
labels = {"method": request.method, "route": _route_label(request),
          "status_code": str(response.status_code)}""")
CODE("""- alert: OpsForgeApiDown
  expr: up{job="opsforge-api"} == 0
  for: 30s
  labels: {severity: critical, service: opsforge-api}""")
P("**Choix.** Le label `route` utilise le *template* FastAPI (`/api/services/{service_id}`), pas l'URL "
  "réelle : sans cela, chaque identifiant créerait une série Prometheus distincte — explosion de "
  "cardinalité (ADR 017). `/metrics` est exclu de son propre comptage. La règle s'appuie sur la "
  "métrique `up` du scrape : elle détecte l'injoignabilité quelle qu'en soit la cause, là où une "
  "métrique applicative disparaîtrait avec l'application ; `for: 30s` est calibré sur le scrape de "
  "15 s, de sorte qu'un raté isolé ne déclenche pas l'alerte.")
P("**Preuve.** Cible `UP`, dashboard alimenté, et cycle `inactive → pending → firing → résolu` observé "
  "en direct (figures 5 à 8, § 3.6). **Limite.** Pas de routage de notification, faute d'Alertmanager.")

H(2, "5.2 Image Docker durcie (CP n° 7)")
P("**Besoin.** Une image reproductible qui ne tourne jamais en root, dans Compose comme dans "
  "Kubernetes.")
CODE("""FROM python:3.12-slim@sha256:c3d81d25b3154142b0b42eb1e61300024426268edeb5b5a26dd7ddf64d9daf28
# [...]
RUN useradd --create-home --shell /usr/sbin/nologin --uid 10001 opsforge \\
    && chown opsforge:opsforge /app
USER opsforge
# [...]
HEALTHCHECK --interval=30s --timeout=5s --start-period=10s --retries=3 \\
  CMD ["python", "-c", "from urllib.request import urlopen; urlopen('http://127.0.0.1:8000/health', timeout=3)"]""")
P("**Choix.** Digest épinglé : deux builds espacés donnent la même base, et le scan porte sur une image "
  "identifiée (ADR 024). L'UID 10001 est réutilisé tel quel dans le securityContext Kubernetes. Le "
  "healthcheck est écrit en Python pur — aucune dépendance ajoutée pour la sonde. **Preuve.** Image "
  "construite à chaque commit en CI ; conteneur `healthy` en Compose ; la même image tourne sous rootfs "
  "en lecture seule dans k3d. **Limite.** Les 19 HIGH et 3 CRITICAL Debian résiduels du scan restent "
  "visibles et sans correctif disponible — assumés.")

H(2, "5.3 Workloads Kubernetes durcis et stockage persistant (CP n° 7)")
P("**Besoin.** Déployer avec le principe du moindre privilège, des démarrages ordonnés, et un stockage "
  "qui survit aux pods.")
CODE("""securityContext:
  runAsNonRoot: true
  runAsUser: 10001
  runAsGroup: 10001
  allowPrivilegeEscalation: false
  readOnlyRootFilesystem: true
  capabilities: {drop: ["ALL"]}
  seccompProfile: {type: RuntimeDefault}
# [...]
readinessProbe: {httpGet: {path: /ready, port: http}}
livenessProbe:  {httpGet: {path: /health, port: http}}""")
P("**Choix.** Les deux probes sont **différentes**, et c'est le point clé : `/ready` (un `SELECT 1` sur "
  "PostgreSQL) retire l'API du Service quand la base est injoignable ; `/health` ne redémarre le "
  "conteneur que si le processus ne répond plus. Utiliser `/ready` en liveness provoquerait des "
  "redémarrages en boucle pendant une panne de base. Un init container `pg_isready` sérialise le "
  "démarrage — l'équivalent du `depends_on` de Compose. PostgreSQL est un **StatefulSet** adossé au PVC "
  "(§ 3.5) ; seul `/tmp` (emptyDir) est inscriptible.")
P("**Preuve.** Vérifié sur cluster réel : pods `1/1 Ready`, pages en 200 sous rootfs en lecture seule, "
  "écriture dans `/app` refusée (audit du 07/08/2026) ; état re-constaté et persistance re-prouvée le "
  "12/08/2026 (§ 3.5). **Limite.** Une réplique de chaque workload — la haute disponibilité n'est pas "
  "l'objectif du périmètre.")

H(2, "5.4 Pipeline d'intégration continue (chaîne 1)")
P("**Besoin.** Vérifier chaque commit : qualité, comportement du domaine, compatibilité avec le moteur "
  "de base réel, image, vulnérabilités.")
CODE("""- name: Lint with Ruff
  run: ruff check .
- name: Run SQLite unit tests
  run: pytest tests/test_app.py
- name: Run PostgreSQL integration test
  env: {DATABASE_URL: postgresql+psycopg://opsforge:opsforge@127.0.0.1:5432/opsforge}
  run: pytest tests/postgres_integration.py
- name: Build Docker image
  run: docker build -t "${IMAGE_NAME}:${IMAGE_TAG}" .
- name: Run Trivy image scan (advisory)
  uses: aquasecurity/trivy-action@v0.36.0
  continue-on-error: true
  with: {scan-type: image, image-ref: "opsforge-api:${{ github.sha }}",
         exit-code: "1", severity: "HIGH,CRITICAL"}  # [...]""")
P("**Choix.** L'ordre est voulu : lint (échec rapide) → 38 tests SQLite (quelques secondes) → test "
  "d'intégration contre un conteneur `postgres:16-alpine` démarré par le job. Ce test crée une base au "
  "nom unique, rejoue le flux complet, puis la supprime : la CI ne peut pas polluer une base de "
  "démonstration. La subtilité Trivy : `exit-code: \"1\"` garde le signal visible quand des HIGH ou "
  "CRITICAL existent, tandis que `continue-on-error: true` maintient le job vert — une politique "
  "advisory explicite (ADR 013), transformable en portail bloquant par une décision documentée.")
P("**Preuve.** Run du candidat final `a9ec694` vérifié le 19/08/2026 via l'API GitHub "
  "(`evidence/github_actions_run.txt`) :")
CODE("""run 32197168814 - CI - branch phase6-operator-ux - sha a9ec694 - conclusion: success

Lint with Ruff ................ success      Build Docker image ............. success
Run SQLite unit tests ......... success      Run Trivy image scan (advisory)  success (*)
Run PostgreSQL integration test success""")
P("(*) L'API GitHub aplatit l'issue d'une étape `continue-on-error` ; l'interface web, elle, "
  "affiche l'état advisory de l'étape.", size=9.5, italic=True)
P("**Limite.** La CI ne publie pas l'image et ne déploie rien : c'est de l'intégration continue, pas "
  "du déploiement continu distant.")

H(2, "5.5 Automatisation du déploiement d'infrastructure avec Ansible (CP n° 2, chaîne 2)")
P("**Besoin.** Remplacer la séquence manuelle documentée (`k3d cluster create` puis une série de "
  "`kubectl apply`) par une commande unique, idempotente et auto-vérifiée.")
P("**Pourquoi Ansible plutôt que Terraform.** Les deux sont cités par le référentiel et légitimes — et "
  "Terraform n'est pas « réservé au cloud ». Terraform est déclaratif et à état : on décrit un état "
  "cible qu'il réconcilie via des providers. Ma tâche est une **orchestration procédurale multi-outils "
  "locale** — préparer le control node, créer le cluster, construire et importer l'image, appliquer des "
  "manifests dans un ordre imposé, attendre, vérifier en HTTP : la forme d'exécution qu'Ansible exprime "
  "naturellement, la collection `kubernetes.core` apportant le déclaratif sur les ressources "
  "Kubernetes. Terraform deviendrait pertinent pour un provisioning cloud, direction de la CP n° 4, non "
  "réalisée (ADR 029).")
CODE("""# role cluster - creation idempotente
- name: Create the k3d cluster
  ansible.builtin.command: >-
    k3d cluster create {{ cluster_name }} --servers 1
    --api-port 127.0.0.1:{{ kubeapi_host_port }}
    --port "{{ api_host_port }}:{{ node_port }}@server:0" --wait
  when: cluster_name not in (k3d_clusters.stdout | from_json | map(attribute='name') | list)""")
CODE("""# role kubernetes_resources - Secret genere, deploiement ordonne et attendu
- name: Create the PostgreSQL Secret from variables
  kubernetes.core.k8s:
    definition:
      kind: Secret
      stringData: {POSTGRES_USER: "{{ db_user }}", POSTGRES_PASSWORD: "{{ db_password }}"}  # [...]
- name: Deploy PostgreSQL and wait until it is ready
  kubernetes.core.k8s: {src: "{{ manifests_path }}/{{ item }}", wait: true}  # [...]""")
CODE("""# role verify - le deploiement se prouve lui-meme
- name: Check /ready through the NodePort (proves PostgreSQL connectivity)
  ansible.builtin.uri:
    url: "http://host.docker.internal:{{ api_host_port }}/ready"
  retries: 12
  until: ready.status is defined and ready.status == 200""")
P("**Choix.** Le cluster n'est créé que s'il n'existe pas, ce qui rend le playbook rejouable sans "
  "danger. **Aucun manifest de Secret n'est versionné** : il est généré au déploiement depuis des "
  "variables (défaut local non sensible, surchargeable par `-e` ou `ansible-vault`). Chaque étage est "
  "appliqué avec `wait: true` — PostgreSQL prêt, puis API prête, puis supervision — et non par un apply "
  "aveugle. Enfin, le run ne réussit **que si l'application répond** : `/ready` prouve la chaîne "
  "complète jusqu'à la base. Le tout s'exécute depuis un **control node conteneurisé** "
  "(`./ansible/run.sh`) aux versions épinglées — indispensable sous Windows, et né d'une vraie "
  "situation de recherche (§ 6).")
P("**Preuve.** Rejouée intégralement le 12/08/2026 sur un cluster isolé "
  "(`evidence/ansible_*.txt`) :")
CODE("""Deploiement complet      : PLAY RECAP  ok=22  changed=9  failed=0
                           "/health -> 200 {'status': 'ok', 'service': 'opsforge'}"
                           "/ready  -> 200 {'status': 'ready', 'service': 'opsforge'}"
Second run (idempotence) : PLAY RECAP  ok=21  changed=2  failed=0
Teardown                 : PLAY RECAP  ok=3   changed=1  failed=0""")
P("**Limites.** Cible k3d locale uniquement ; le control node conteneurisé est la seule méthode "
  "validée ; `validate_certs: false` est limité à la connexion de contrôle locale vers le cluster "
  "éphémère.")

H(2, "5.6 Sauvegarde et restauration sûres")
P("**Besoin.** Prouver qu'une sauvegarde existe **et** qu'elle est restaurable, sans jamais risquer la "
  "base par une fausse manœuvre.")
CODE("""if ($MainDatabase) {
    $confirmation = Read-Host "This will replace objects in database '$dbName'. Type RESTORE to continue"
    if ($confirmation -cne "RESTORE") { Write-Host "Restore cancelled." ; return }
}""")
P("**Choix.** Format custom, inspectable par `pg_restore --list` ; restauration **par défaut dans une "
  "base temporaire** de vérification, avec comptage des tables puis nettoyage ; la base principale "
  "exige le drapeau `-MainDatabase` **et** la saisie exacte de `RESTORE`, sensible à la casse — une "
  "touche Entrée pressée trop vite annule l'opération ; l'API est arrêtée pendant une restauration "
  "réelle, puis redémarrée. **Preuve.** Rejouée le 12/08/2026 : archive de 24 543 octets créée, puis "
  "restauration vérifiée sur 6 tables dans `opsforge_restore_verify` (§ 3.7). **Limite.** Locale, non "
  "planifiée, non chiffrée.")

# =============================================================== PARTIE 6 ==
H(1, "6. Situation de travail ayant nécessité une recherche", break_before=True)
H(2, "Le fichier ansible.cfg silencieusement ignoré dans le control node conteneurisé")
P("**Le problème.** Situation survenue lors de l'intégration d'Ansible (10 et 11 août 2026), traçable "
  "dans le dépôt (commits `9d04fc4`, `0b21505`, `732d6fa`, `e808ccf`). Ansible ne s'exécutant pas "
  "nativement sous Windows, j'avais choisi un control node conteneurisé : une image Docker embarquant "
  "Ansible et les CLI, le dépôt monté en `/work`. Première exécution du playbook dans ce conteneur :")
CODE("""skipping: no hosts matched""")
P("Aucun hôte trouvé — alors que l'inventaire (`localhost ansible_connection=local`) existait, déclaré "
  "dans `ansible.cfg` juste à côté du playbook.")

P("**Le diagnostic.** Hypothèses testées dans l'ordre. Inventaire mal écrit ? Non : il fonctionnait "
  "passé à la main. Mauvais répertoire de travail ? Non : le `WORKDIR` était correct. Le fichier "
  "`ansible.cfg` est-il seulement lu ? `ansible --version` affichait `config file = None` : Ansible "
  "ignorait silencieusement la configuration, donc l'inventaire qu'elle déclare.")

P("**La recherche.** La documentation officielle d'Ansible décrit un comportement de sécurité précis : "
  "**un `ansible.cfg` situé dans un répertoire courant inscriptible par tous (*world-writable*) est "
  "refusé**, pour empêcher l'injection d'une configuration malveillante dans un répertoire partagé ; un "
  "chemin désigné via `ANSIBLE_CONFIG` reste honoré. Or c'était exactement mon contexte, sans que je "
  "l'aie provoqué : un dépôt Windows monté par bind-mount Docker apparaît world-writable côté Linux. Le "
  "comportement d'Ansible était correct — c'est mon environnement qui le déclenchait.")

P("**Les options et le choix.** Trois pistes : `chmod` du répertoire à chaque run (fragile) ; tout "
  "passer en ligne de commande sans `ansible.cfg` (ce qui disperse la configuration) ; fixer "
  "`ANSIBLE_CONFIG` dans l'image — la solution prévue par l'outil. J'ai retenu la troisième, doublée "
  "d'une ceinture de sécurité : `ENV ANSIBLE_CONFIG=/work/ansible/ansible.cfg` cuit dans l'image, "
  "**plus** un `-i inventory.ini` explicite dans `run.sh`.")

P("**Le second enseignement — le plus important.** En corrigeant, j'ai découvert plus embarrassant que "
  "le bug : mes validations initiales avaient été faites avec des commandes ajustées à la main pendant "
  "le débogage, et le `run.sh` documenté ne portait pas ces réglages — ni l'inventaire explicite, "
  "ni `--add-host=host.docker.internal:host-gateway`, ni `MSYS_NO_PATHCONV=1`. Quiconque aurait "
  "exécuté le `./run.sh` documenté — le jury, par exemple — aurait reproduit l'échec initial. Le "
  "commit de correction le dit sans détour : *« run.sh did not carry the settings the successful runs "
  "actually used »*. J'ai donc :",
  space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
BUL([
    "aligné `run.sh` sur le chemin réellement validé ;",
    "**re-testé l'intégralité du parcours en n'utilisant que `./run.sh`** sur un cluster isolé "
    "(déploiement complet, second run idempotent, teardown), puis épinglé les versions exactes du "
    "control node validé pour garder la validation reproductible ;",
    "retiré de la documentation une affirmation que je ne pouvais pas prouver : le commit initial "
    "prétendait que le playbook « tourne aussi directement sur un hôte Linux/WSL », un chemin jamais "
    "validé qui n'aurait pas fonctionné tel quel (le kubeconfig est réécrit vers "
    "`host.docker.internal`, un nom fourni par le conteneur). Le control node est devenu "
    "**la seule méthode supportée et validée**.",
])
CALLOUT("**Ce que j'en retiens.** Un comportement de sécurité d'un outil peut n'apparaître que dans un "
        "contexte d'exécution particulier : le diagnostic passe par la documentation, pas par des "
        "essais au hasard. Et surtout, **l'artefact documenté doit être exactement celui qui a été "
        "validé** — retirer de sa propre documentation une affirmation non prouvée est une correction "
        "au même titre qu'un correctif de code.")

# =============================================================== PARTIE 7 ==
H(1, "7. Synthèse des validations", break_before=True)
TABLE(
    ["Domaine", "Preuve principale", "Résultat"],
    [
        ["Application et tests",
         "38 tests SQLite + 1 test d'intégration PostgreSQL (base éphémère)",
         "Verts en local et en CI"],
        ["Intégration continue",
         "Run GitHub Actions du candidat final `a9ec694`",
         "`success` (toutes étapes) — vérifié le 19/08/2026"],
        ["Conteneurs et Kubernetes",
         "Pods `1/1 Running`, PVC `Bound`, durcissement vérifié sous rootfs en lecture seule",
         "Constaté le 07/08 et re-constaté le 12/08/2026"],
        ["Persistance",
         "Donnée survivant à la destruction puis à la recréation du pod PostgreSQL (UID différent)",
         "Prouvée en phase 4, **re-prouvée le 12/08/2026**"],
        ["Supervision",
         "Cible UP, dashboard alimenté, cycle d'alerte `inactive → firing → résolu`",
         "Validée en phase 5, **rejouée en direct le 12/08/2026**"],
        ["Automatisation d'infrastructure",
         "Déploiement complet en une commande, idempotence, teardown, `/health` et `/ready` en 200",
         "Validée aux commits Ansible, **rejouée le 12/08/2026** (`ok=22/9`, `ok=21/2`)"],
        ["Sauvegarde",
         "Archive produite et restauration vérifiée en base temporaire (6 tables)",
         "Validée en phase 3, **rejouée le 12/08/2026**"],
    ],
    widths=[3.6, 7.4, 5.6],
    first_col_bold=True,
)
P("Les preuves du 12/08/2026 — logs bruts et captures — sont versionnées sous "
  "`deliverables/evidence/` et `deliverables/assets/` ; l'historique détaillé phase par phase reste "
  "dans `docs/PHASE<i>_VERIFICATION.md` (chronologie en annexe A).")

H(2, "Clôture de la phase 6 : ce que la revue humaine a trouvé")
P("La revue manuelle — parcours desktop, workflow opérateur complet, comportement responsive — a été "
  "déroulée du 12 au 19/08/2026, guidée par un runbook dédié. Elle a détecté **deux défauts réels, "
  "invisibles des 35 tests automatisés alors en place** : une erreur HTTP 422 sur les formulaires de "
  "filtres lorsque aucun service n'était sélectionné, et une table Incidents illisible sur mobile. Les "
  "deux ont été corrigés en branches dédiées, couverts par des tests de non-régression qui ont porté la "
  "suite à **38 tests**, puis revalidés humainement sur le candidat fusionné ; la phase 6 a alors été "
  "**explicitement validée le 19/08/2026** (PR #4 à #6, candidat final `a9ec694`).")
CALLOUT("Cette séquence est un enseignement en soi : les deux défauts passaient sous le radar de tests "
        "qui vérifient des réponses HTTP et des structures de page, mais ne simulent ni la soumission "
        "réelle d'un formulaire ni un rendu à 390 px. **Une procédure de test manuelle documentée n'est "
        "pas un reliquat d'un projet non automatisé : c'est la couche qui attrape ce que "
        "l'automatisation ne voit pas.**")

# =============================================================== PARTIE 8 ==
H(1, "8. Limites assumées et évolutions", break_before=True)
P("Toutes ces limites sont documentées dans le dépôt (`docs/RISKS_AND_TECHNICAL_DEBT.md`, ADR). Ce sont "
  "des choix de périmètre, pas des fonctions prétendues :")
TABLE(
    ["Domaine", "Limite", "Évolution naturelle"],
    [
        ["Infrastructure",
         "k3d mono-nœud local, pas de cloud ; stockage `local-path` local au nœud, non distribué",
         "Kubernetes managé + Terraform (CP n° 4) ; stockage réseau"],
        ["Livraison",
         "Pas de registre, pas de CD distant ; import k3d (manuel ou par Ansible)",
         "Registre + déploiement déclenché par la CI"],
        ["Application",
         "Pas d'authentification (acteurs déclaratifs) ; unicité d'incident actif garantie par "
         "l'application (409), pas par contrainte en base",
         "Authentification ; contrainte PostgreSQL partielle"],
        ["Schéma",
         "`create_all()` et pont additif, pas d'historique de migrations",
         "Alembic"],
        ["Supervision",
         "Pas d'Alertmanager ni de notifications ; stockage éphémère ; Grafana `admin/admin` en local ; "
         "accès par port-forward ; métriques techniques ; **états métier simulés**",
         "Alertmanager, persistance, métriques métier, ingestion des alertes dans OpsForge"],
        ["Sécurité",
         "Trivy advisory (19 HIGH / 3 CRITICAL sans correctif, visibles, non bloquants) ; pas de TLS",
         "Seuil de blocage explicite ; TLS"],
        ["Sauvegardes",
         "Locales, non planifiées, non chiffrées, sans rotation",
         "Planification, chiffrement, externalisation"],
        ["Outillage",
         "Écart de versions entre le kubectl du control node (v1.31) et k3s (v1.35)",
         "Alignement des versions épinglées"],
        ["Contexte",
         "Projet individuel : le critère relationnel « échanges avec les développeurs » (CP n° 10) est "
         "porté par des décisions tracées, pas par une équipe",
         "—"],
    ],
    widths=[2.9, 8.0, 5.7],
    first_col_bold=True,
)

# ============================================================== CONCLUSION ==
H(1, "Conclusion", break_before=True)
P("**Ce que le projet démontre.** Les trois compétences obligatoires, chacune avec une preuve "
  "distincte, rejouable et rejouée : une infrastructure locale complète déployée et vérifiée par "
  "Ansible en une commande idempotente ; des conteneurs durcis et orchestrés dans deux environnements, "
  "avec une persistance prouvée plutôt que supposée ; une supervision réelle dont l'alerte s'est "
  "déclenchée puis résolue lors d'une panne provoquée. Autour : 38 tests unitaires et un test "
  "d'intégration PostgreSQL, une CI avec scan de vulnérabilités, des sauvegardes restaurables, "
  "29 décisions d'architecture documentées.")
P("**Mes satisfactions.** La méthode — six phases finies, validées et datées une à une — qui a permis "
  "au projet de survivre sans dégât à un changement de poste de travail. La correction de trajectoire "
  "de la fin de projet : confronter le travail aux critères exacts du référentiel, constater qu'un "
  "déploiement documenté mais manuel ne prouvait pas la compétence d'automatisation, et fermer l'écart "
  "proprement, avant l'examen. Et une exigence tenue de bout en bout : distinguer partout ce qui a été "
  "testé de ce qui a seulement été écrit.")
P("**Mes difficultés.** Le débogage du control node Ansible (§ 6), le plus formateur — avec sa leçon : "
  "l'artefact documenté doit être exactement celui qui a été validé. La revue visuelle finale, qui n'a "
  "pas pu être automatisée : déroulée manuellement, elle a révélé deux défauts réels que les 35 tests "
  "d'alors ne voyaient pas — preuve que le contrôle humain reste une étape à part entière. Et, en "
  "continu, "
  "tenir le périmètre : dire non à tout ce qui aurait grossi le projet sans le rendre plus défendable.")
P("Le projet est gelé au commit `a9ec694` : c'est cet état, reproductible et documenté, que je présente "
  "au jury.")

# ================================================================= ANNEXES ==
H(1, "Annexe A — Chronologie détaillée du projet", break_before=True)
TABLE(
    ["Date", "Événement", "Commits"],
    [
        ["16/06/2026", "Validation du MVP (7 tests ; travail pré-Git)", "—"],
        ["17/06/2026", "Premier commit : application MVP et workflow CI", "`ad9b9df`"],
        ["18/06/2026",
         "Phase 2 validée sur un run GitHub Actions vert ; gouvernance (index, risques, protocole)",
         "`9e0666f`, `917378b`"],
        ["06/07/2026", "Phase 3 : sauvegarde et restauration, workflow Git", "`a317969`, `34f785b`"],
        ["07-09/07/2026",
         "Phase 4 : k3d + PostgreSQL/PVC (4A, vérifiée le 07), API (4B), preuve de persistance, "
         "validation",
         "`c386c84`, `2772b50`"],
        ["09-14/07/2026",
         "Phase 5 : métriques (5A), Prometheus (5B), Grafana (5C), alerte et panne provoquée (5D), "
         "validation",
         "`84fb228` → `23194f0`, `9bcb271`"],
        ["15-17/07/2026",
         "Phase 6 : candidat produit opérateur — débuté sur `main`, poursuivi sur la première branche "
         "dédiée `phase6-operator-ux`",
         "`83469cb` → `230d07a`"],
        ["07/08/2026",
         "Audit : branche d'expérimentation revue (jamais fusionnée, conservée comme trace), "
         "ré-implémentation propre en 6 commits sur `integration/phase6-audit` (35 tests, durcissement "
         "Kubernetes vérifié sur cluster réel)",
         "`fb4d77c` → `fd3dca9`"],
        ["10-11/08/2026",
         "Ansible (CP n° 2) : implémentation, corrections de la situation de recherche (§ 6), "
         "épinglage, honnêteté documentaire",
         "`9d04fc4`, `0b21505`, `732d6fa`, `e808ccf`"],
        ["11/08/2026",
         "Intégration finale par Pull Requests avec commits de merge : PR #1 (audit), PR #2 (Ansible), "
         "PR #3 (synchronisation documentaire) → candidat gelé",
         "`489552f`, `0becdf8`, `8ab0f70`"],
        ["12/08/2026",
         "Production des preuves du présent dossier : re-déploiement Ansible complet, idempotence et "
         "teardown, persistance re-prouvée, cycle d'alerte rejoué en direct, sauvegarde et restauration "
         "rejouées, CI du candidat vérifiée, captures d'écran",
         "branche `jury/dossier-fil-rouge`"],
        ["12-19/08/2026",
         "Revue manuelle finale (runbook dédié) : deux défauts réels détectés — erreur 422 des "
         "formulaires de filtres et table Incidents illisible en mobile — corrigés en branches dédiées "
         "avec tests de non-régression (35 → 38 tests), plus le favicon ; revalidation humaine sur le "
         "candidat corrigé",
         "PR #4 `3fc7707`, PR #5 `f6e4a79`"],
        ["19/08/2026",
         "**Phase 6 validée explicitement** ; clôture documentaire (roadmap, vérification, risques) → "
         "candidat final",
         "PR #6 `a9ec694`"],
    ],
    widths=[2.5, 9.9, 4.2],
    first_col_bold=True,
    font=9.0,
)
P("Le candidat d'examen `phase6-operator-ux @ a9ec694` totalise **45 commits**, dont 6 commits de "
  "merge. La branche `main` porte l'état des phases 1 à 6 initiales.")

H(1, "Annexe B — Inventaire des preuves", break_before=True)
H(2, "B.1 Captures et schémas (versionnés sous deliverables/assets/)")
TABLE(
    ["Fichier", "Contenu", "Usage"],
    [
        ["`diagrams/schema_architecture.png`", "Architecture logique et les deux chaînes d'outillage",
         "Figure 2"],
        ["`diagrams/schema_topologie.png`", "Topologie de déploiement sur le poste", "Figure 3"],
        ["`screenshots/01_overview.png`", "Vue d'ensemble de la console", "Figure 1"],
        ["`screenshots/02_alerts.png`", "File d'alertes (états, filtres)", "Réserve orale"],
        ["`screenshots/03_incident_command_center.png`",
         "Command Center : contexte, runbooks, timeline, exécutions", "Figure 4"],
        ["`screenshots/04_activity.png`", "Journal d'audit global", "Réserve orale"],
        ["`screenshots/05_monitoring.png`", "Page Monitoring : réel contre simulé", "Figure 5"],
        ["`screenshots/06_prometheus_targets_up.png`", "Cible Prometheus `opsforge-api (1/1 up)`",
         "Figure 6 (*)"],
        ["`screenshots/07_prometheus_alert_firing.png`", "`OpsForgeApiDown` en FIRING", "Figure 7 (*)"],
        ["`screenshots/08_grafana_dashboard.png`", "Dashboard « OpsForge Monitoring » alimenté",
         "Figure 8"],
        ["`screenshots/09_command_center_mobile.png`", "Command Center en mobile (~390 px)",
         "Réserve orale"],
        ["`screenshots/10_incidents_mobile.png`",
         "Table Incidents en mobile après correction (défaut trouvé en revue manuelle)",
         "Réserve orale"],
    ],
    widths=[6.3, 7.3, 3.0],
    font=9.0,
)

P("(*) Les figures 6 et 7 sont recadrées sur la zone utile de la capture ; les fichiers ci-dessus sont "
  "les captures complètes, non retouchées. Les versions recadrées utilisées dans le document sont "
  "versionnées sous `deliverables/assets/figures/`.", size=9.2, italic=True)

H(2, "B.2 Journaux d'exécution bruts (versionnés sous deliverables/evidence/)")
TABLE(
    ["Fichier", "Contenu"],
    [
        ["`ansible_fresh_deploy.txt`",
         "Déploiement complet par Ansible : `ok=22 changed=9`, `/health` et `/ready` en 200"],
        ["`ansible_second_run.txt`", "Second run : idempotence `ok=21 changed=2`"],
        ["`ansible_teardown.txt`", "Suppression du cluster de test"],
        ["`kubernetes_state.txt`", "État constaté : pods, services, PVC, nœud"],
        ["`pvc_persistence.txt`", "Preuve de persistance : marqueur et UID du pod avant/après"],
        ["`prometheus_alert_cycle.txt`", "Cycle d'alerte complet horodaté et payload de l'alerte"],
        ["`backup_restore.txt`", "Archive produite et restauration vérifiée"],
        ["`github_actions_run.txt`", "Run CI du candidat final `a9ec694`, étape par étape"],
    ],
    widths=[5.2, 11.4],
    font=9.0,
)
P("Les environnements de production de ces preuves sont tracés : les captures d'interface ont été "
  "prises sur un environnement Docker Compose isolé et éphémère, et les preuves Kubernetes, Ansible et "
  "supervision sur un cluster k3d jetable, distinct du cluster de travail.")

H(1, "Annexe C — Glossaire", break_before=True)
TABLE(
    ["Terme", "Définition dans le contexte du projet"],
    [
        ["ADR", "*Architecture Decision Record* : décision consignée (contexte, décision, raison, "
                "conséquences) — 29 dans `docs/DECISIONS.md`"],
        ["Advisory (scan)", "Résultats visibles mais non bloquants pour le pipeline"],
        ["Command Center", "Vue de travail d'un incident : contexte, responsable, transitions, "
                           "runbooks, exécutions, timeline"],
        ["Control node", "Machine — ici un conteneur — depuis laquelle Ansible s'exécute"],
        ["CP / REAC", "Compétence professionnelle / Référentiel Emploi Activités Compétences du titre"],
        ["DoD", "*Definition of Done* : conditions vérifiables de fin d'une phase"],
        ["Idempotence", "Une automatisation rejouée converge sans rien recréer ni casser"],
        ["k3d", "Cluster Kubernetes k3s exécuté dans des conteneurs Docker"],
        ["Liveness / readiness",
         "« Le processus vit-il ? » (redémarrage) / « Peut-il servir ? » (retrait du Service)"],
        ["NodePort", "Service Kubernetes exposant un port du nœud (30080 → `127.0.0.1:8080`)"],
        ["PVC", "*PersistentVolumeClaim* : demande de stockage persistant (1 Gi, `local-path`)"],
        ["Runbook", "Procédure : checklist manuelle ou action automatisée approuvée dans le code"],
        ["Scrape", "Collecte périodique des métriques par Prometheus (toutes les 15 s)"],
        ["StatefulSet",
         "Contrôleur pour workloads à état : identité et stockage stables (PostgreSQL)"],
    ],
    widths=[3.6, 13.0],
    first_col_bold=True,
    font=9.4,
)

cp = doc.core_properties
cp.title = "Dossier de projet OpsForge"
cp.subject = "Titre professionnel Administrateur systeme DevOps (TP-01414 / RNCP 36061)"
cp.author = "Dyllan Thouvignon"
cp.category = "Dossier de projet"
cp.comments = "Candidat technique gele : phase6-operator-ux @ a9ec694"
cp.language = "fr-FR"

doc.save(OUT)
print("DOCX ecrit :", OUT)
